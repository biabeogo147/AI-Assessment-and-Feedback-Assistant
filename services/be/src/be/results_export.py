"""Bảng điểm của một đề trong một lớp, dựng thành tệp `.xlsx` (việc 20, plan `results-xlsx-export`).

Một trang tính, một dòng đầu, rồi mỗi học sinh của lớp một dòng theo mã học sinh: *Mã học sinh*,
*Họ tên*, *Câu 1*…*Câu N*, *Tổng*. Bảng phẳng để dán thẳng lên sổ điểm; khớp mẫu vnEdu/SMAS khi có
tệp mẫu thật.

Em chưa nộp bài vẫn có dòng, và các ô điểm **để trống**: trống và 0 là hai sự thật khác nhau -- 0 là
em ấy làm bài và sai hết (Decision 1 của plan).

Hàm ở đây thuần: nhận dữ liệu đã đọc, trả byte. Route và quyền (chỉ tác giả, ADR-22) ở
`teacher_reads.results_xlsx`.
"""

import re
import unicodedata
from collections.abc import Sequence
from dataclasses import dataclass
from io import BytesIO
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Excel giới hạn tên trang tính ở 31 ký tự và cấm vài ký tự.
_SHEET_NAME_LIMIT = 31
_SHEET_FORBIDDEN = str.maketrans({one: " " for one in "[]:*?/\\"})

# Chuỗi mở đầu bằng một trong các ký tự này thì Excel đọc như công thức (chèn công thức, CSV/XLSX
# injection): một họ tên *"=HYPERLINK(...)"* thành một link. Ô ấy được ghi là chữ, có dấu nháy ẩn.
_FORMULA_START = ("=", "+", "-", "@")


def _as_text(cells) -> None:
    """Ghi các ô chữ mở đầu như công thức thành chữ thật, có dấu nháy ẩn."""
    for cell in cells:
        if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_START):
            cell.data_type = "s"
            cell.quotePrefix = True


def attachment(name: str) -> dict[str, str]:
    """Header `Content-Disposition` để tải một tệp tên tiếng Việt.

    `filename` ASCII cho trình duyệt cũ, `filename*` UTF-8 cho tên tiếng Việt (RFC 6266, RFC 5987).
    Bản ASCII bỏ dấu rồi thay mọi ký tự lạ -- một `"` hay một dấu xuống dòng trong tên đề phá cả
    header (review pha 1 của việc 20).

    Args:
        name: Tên tệp, có đuôi.

    Returns:
        Một header, để trộn vào `Response(headers=...)`.
    """
    plain = re.sub(
        r"[^A-Za-z0-9 ._-]",
        "_",
        unicodedata.normalize("NFKD", name.replace("đ", "d").replace("Đ", "D"))
        .encode("ascii", "ignore")
        .decode(),
    )
    return {
        "Content-Disposition": (
            f"attachment; filename=\"{plain}\"; filename*=UTF-8''{quote(name, safe='')}"
        )
    }


@dataclass(frozen=True)
class ResultLine:
    """Một học sinh trong bảng điểm.

    Attributes:
        student_code: Mã học sinh.
        full_name: Họ tên.
        marks: Điểm từng câu theo thứ tự câu, hay `None` khi em chưa nộp. Một câu không có
            dòng sổ điểm là `None` ở ô của câu ấy: trống, không phải 0.
    """

    student_code: str
    full_name: str
    marks: Sequence[float | None] | None


def build_workbook(class_name: str, question_count: int, lines: Sequence[ResultLine]) -> bytes:
    """Dựng tệp `.xlsx` của một bảng điểm.

    Args:
        class_name: Tên lớp -- tên trang tính.
        question_count: Số câu của đề, tức số cột *Câu k*.
        lines: Mỗi học sinh một dòng, đã sắp.

    Returns:
        Byte của tệp.
    """
    book = Workbook()
    sheet = book.active
    # Dấu nháy ở hai đầu tên làm Excel đòi *sửa tệp*.
    sheet.title = (class_name.translate(_SHEET_FORBIDDEN).strip(" '") or "Bảng điểm")[
        :_SHEET_NAME_LIMIT
    ]
    head = ["Mã học sinh", "Họ tên", *(f"Câu {n}" for n in range(1, question_count + 1)), "Tổng"]
    sheet.append(head)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for line in lines:
        if line.marks is None:
            sheet.append([line.student_code, line.full_name])
        else:
            total = round(sum(one for one in line.marks if one is not None), 2)
            sheet.append([line.student_code, line.full_name, *line.marks, total])
        _as_text(sheet[sheet.max_row][:2])
    sheet.freeze_panes = "C2"
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 28
    for index in range(3, question_count + 4):
        sheet.column_dimensions[get_column_letter(index)].width = 8
    out = BytesIO()
    book.save(out)
    return out.getvalue()


@dataclass(frozen=True)
class PasswordLine:
    """Một học sinh trong tệp mật khẩu (việc 22, plan `sign-in-and-accounts`).

    Attributes:
        full_name: Họ tên.
        username: Tên đăng nhập.
        password: Mật khẩu rõ -- chỉ đi qua đây, không bao giờ được cất (Decision 3 của plan).
    """

    full_name: str
    username: str
    password: str


def build_password_workbook(class_name: str, lines: Sequence[PasswordLine]) -> bytes:
    """Dựng tệp `.xlsx` mật khẩu của những tài khoản vừa tạo, để giáo viên in phát cho lớp.

    Args:
        class_name: Tên lớp, thành tên trang tính.
        lines: Các dòng, theo thứ tự hiện trên màn 28c.

    Returns:
        Byte của tệp. Không ghi gì xuống đĩa hay database.
    """
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = (class_name.translate(_SHEET_FORBIDDEN).strip(" '") or "Mật khẩu")[
        :_SHEET_NAME_LIMIT
    ]
    sheet.append(["Họ và tên", "Tên đăng nhập", "Mật khẩu"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for line in lines:
        sheet.append([line.full_name, line.username, line.password])
        _as_text(sheet[sheet.max_row])
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 16
    out = BytesIO()
    book.save(out)
    return out.getvalue()
