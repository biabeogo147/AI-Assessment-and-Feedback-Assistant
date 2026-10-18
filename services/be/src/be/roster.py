"""Thêm học sinh vào lớp và đặt lại mật khẩu (việc 22, plan `sign-in-and-accounts`, ADR-36).

Ba bước của Figma 28 / 28b / 28c, mỗi bước một route:
1. **xem trước** -- từ các tên gõ tay, hay từ tệp `.csv`/`.xlsx` một cột họ tên. Không ghi gì: trả
   tên đăng nhập sẽ cấp và đánh dấu dòng *trùng tên trong lớp*;
2. **tạo tài khoản** -- dòng trùng chỉ được tạo khi giáo viên đã đánh dấu *Xác nhận*. BE tính lại
   mọi thứ từ đầu, không tin tên đăng nhập hay ghi chú FE gửi về;
3. **tệp mật khẩu** -- dựng từ chính các dòng FE gửi lên, không lưu gì (Decision 3 của plan).

Mật khẩu rõ chỉ xuất hiện trong phản hồi của bước 2 và của *Đặt lại mật khẩu*, một lần. Cột
`password_hash` giữ `scrypt` của nó.
"""

import asyncio
import csv
import io
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy import delete
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from be.accounts import hashed, new_password, next_student_codes, plain, usernames_for
from be.db import get_session
from be.identity import Asking, current_teacher
from be.results_export import PasswordLine, attachment, build_password_workbook
from be.teacher_reads import _owned_class
from schema.models import AccountSession, Student, Teacher

router = APIRouter(prefix="/api", tags=["teacher-roster"])

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Một lớp ở Việt Nam dưới 50 em; 200 đỡ cả một lần nhập cho bốn lớp ghép mà vẫn chặn một tệp lạ.
MOST_ROWS = 200
MOST_FILE_BYTES = 1_000_000
MOST_NAME = 128
# Đọc tới đây thì thôi: quá `MOST_ROWS` tên là bị từ chối dù sao, và một tệp 1 MB có thể chứa
# hàng trăm nghìn dòng trống -- dựng hết thành list trước khi đếm là chặn BE vài giây (review).
_READ_ROWS = MOST_ROWS * 5
# Mật khẩu rõ không được nằm lại trong cache của trình duyệt hay proxy nào.
_NO_STORE = {"Cache-Control": "no-store"}
# Dòng đầu của tệp là tiêu đề khi nó là một trong các chữ này (bỏ dấu, chữ thường).
_HEADINGS = {"hovaten", "hoten", "ten", "hotenhocsinh", "name", "fullname"}

Note = Literal["new", "duplicate"]


class NamesIn(BaseModel):
    """Các tên giáo viên gõ ở tab *Gõ tên* (Figma 28), mỗi dòng một tên."""

    names: list[str]


class PreviewRow(BaseModel):
    """Một dòng của bảng xem trước (Figma 28b).

    Attributes:
        full_name: Họ tên đã dọn khoảng trắng.
        username: Tên đăng nhập **sẽ** cấp nếu tạo ngay bây giờ.
        note: `new`, hay `duplicate` khi lớp đã có -- hoặc loạt này đã có ở dòng trên -- một em
            cùng họ tên.
    """

    full_name: str
    username: str
    note: Note


class Preview(BaseModel):
    """Bảng xem trước.

    Attributes:
        rows: Các dòng theo thứ tự nhập.
    """

    rows: list[PreviewRow]


class CreateRow(BaseModel):
    """Một dòng giáo viên gửi lên ở bước tạo.

    Attributes:
        full_name: Họ tên.
        confirmed: Ô *Xác nhận* -- chỉ có nghĩa ở dòng trùng.
    """

    full_name: str
    confirmed: bool = False


class CreateIn(BaseModel):
    rows: list[CreateRow]


class Account(BaseModel):
    """Một tài khoản vừa tạo hay vừa đặt lại mật khẩu (Figma 28c, 10b).

    Attributes:
        student_id: Id.
        full_name: Họ tên.
        student_code: Mã học sinh do hệ thống cấp.
        username: Tên đăng nhập.
        password: Mật khẩu rõ -- **chỉ lần này**. BE không giữ bản nào để trả lại.
    """

    student_id: str
    full_name: str
    student_code: str
    username: str
    password: str


class Created(BaseModel):
    """Kết quả bước tạo.

    Attributes:
        accounts: Các tài khoản vừa tạo, theo thứ tự nhập.
        skipped: Số dòng trùng chưa xác nhận nên không tạo.
    """

    accounts: list[Account]
    skipped: int


class PasswordSheetIn(BaseModel):
    """Các dòng của màn 28c, để dựng tệp mật khẩu."""

    rows: list[PasswordLine]


def _clean(name: str) -> str:
    return " ".join(name.split())


def _key(name: str) -> str:
    """Khoá so trùng: khoảng trắng gộp, không phân biệt hoa thường. Giữ dấu -- *Hà* khác *Hạ*."""
    return _clean(name).casefold()


def _names(raw: list[str]) -> list[str]:
    """Dọn một loạt tên: bỏ dòng trống, gộp khoảng trắng, chặn quá dài hay quá nhiều.

    Raises:
        HTTPException: 422 khi không còn tên nào, quá `MOST_ROWS` tên, hay một tên quá dài.
    """
    names = [_clean(one) for one in raw]
    names = [one for one in names if one]
    if not names:
        raise HTTPException(status_code=422, detail="Danh sách chưa có họ tên nào.")
    if len(names) > MOST_ROWS:
        raise HTTPException(
            status_code=422, detail=f"Mỗi lần thêm được tối đa {MOST_ROWS} học sinh."
        )
    for one in names:
        if len(one) > MOST_NAME:
            raise HTTPException(status_code=422, detail=f"Họ tên quá dài: “{one[:40]}…”.")
    return names


def _cells(filename: str, data: bytes) -> list[str]:
    """Cột đầu tiên có chữ của mỗi dòng trong tệp; dòng tiêu đề bị bỏ. Đọc tối đa `_READ_ROWS` dòng.

    Hàm đồng bộ và tốn CPU: route gọi nó qua `asyncio.to_thread`.

    Raises:
        HTTPException: 422 khi tệp không phải `.csv`/`.xlsx` hay đọc không được.
    """
    lower = filename.lower()
    rows: list[list[object]] = []
    if lower.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=422, detail="Tệp .csv cần lưu ở mã UTF-8.") from exc
        try:
            for row in csv.reader(io.StringIO(text)):
                rows.append(list(row))
                if len(rows) >= _READ_ROWS:
                    break
        except csv.Error as exc:  # một ô quá dài, hay dấu nháy không đóng
            raise HTTPException(status_code=422, detail="Không đọc được tệp .csv này.") from exc
    elif lower.endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                for row in book.worksheets[0].iter_rows(values_only=True):
                    rows.append(list(row))
                    if len(rows) >= _READ_ROWS:
                        break
            finally:
                book.close()
        except Exception as exc:  # openpyxl ném đủ loại lỗi với một tệp hỏng, lúc mở lẫn lúc đọc
            raise HTTPException(status_code=422, detail="Không đọc được tệp .xlsx này.") from exc
    else:
        raise HTTPException(status_code=422, detail="Chỉ nhận tệp .csv hoặc .xlsx.")

    names = [
        next((str(cell).strip() for cell in row if cell is not None and str(cell).strip()), "")
        for row in rows
    ]
    head = next((index for index, one in enumerate(names) if one), None)
    if head is not None and plain(names[head]) in _HEADINGS:
        names[head] = ""
    return names


async def _preview(session: AsyncSession, existing: list[Student], names: list[str]) -> Preview:
    seen = {_key(one.full_name) for one in existing}
    notes: list[Note] = []
    for name in names:
        notes.append("duplicate" if _key(name) in seen else "new")
        seen.add(_key(name))
    usernames = await usernames_for(session, names)
    return Preview(
        rows=[
            PreviewRow(full_name=name, username=username, note=note)
            for name, username, note in zip(names, usernames, notes, strict=True)
        ]
    )


@router.post("/teacher/classes/{class_id}/roster/preview", response_model=Preview)
async def preview_names(
    class_id: str,
    body: NamesIn,
    teacher: Teacher = Depends(current_teacher),
    session: AsyncSession = Depends(get_session),
) -> Preview:
    """Bước 1 từ tab *Gõ tên*. Không ghi gì.

    Raises:
        HTTPException: 404 khi lớp không phải của giáo viên này; 422 khi danh sách rỗng hay quá dài.
    """
    school_class = await _owned_class(session, Asking.of(teacher), class_id)
    return await _preview(session, list(school_class.students), _names(body.names))


@router.post("/teacher/classes/{class_id}/roster/preview-file", response_model=Preview)
async def preview_file(
    class_id: str,
    file: UploadFile = File(...),
    teacher: Teacher = Depends(current_teacher),
    session: AsyncSession = Depends(get_session),
) -> Preview:
    """Bước 1 từ tab *Tải tệp .csv / .xlsx*: một cột họ tên, dòng tiêu đề nếu có thì bỏ.

    Raises:
        HTTPException: 404 khi lớp không phải của giáo viên này; 413 khi tệp quá lớn; 422 khi tệp
            sai loại, hỏng, hay không có tên nào.
    """
    school_class = await _owned_class(session, Asking.of(teacher), class_id)
    data = await file.read(MOST_FILE_BYTES + 1)
    if len(data) > MOST_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Tệp danh sách lớn hơn 1 MB.")
    names = _names(await asyncio.to_thread(_cells, file.filename or "", data))
    return await _preview(session, list(school_class.students), names)


@router.post("/teacher/classes/{class_id}/roster", response_model=Created)
async def create_accounts(
    class_id: str,
    body: CreateIn,
    response: Response,
    teacher: Teacher = Depends(current_teacher),
    session: AsyncSession = Depends(get_session),
) -> Created:
    """Bước 2: tạo tài khoản. Dòng trùng chưa xác nhận thì bỏ qua, không lỗi.

    Raises:
        HTTPException: 404 khi lớp không phải của giáo viên này; 409 khi một tên đăng nhập hay mã
            vừa bị lấy mất giữa chừng -- giáo viên bấm lại là được.

    Side effects:
        Thêm các dòng `students` rồi commit.
    """
    school_class = await _owned_class(session, Asking.of(teacher), class_id)
    names = _names([row.full_name for row in body.rows])
    confirmed = [row.confirmed for row in body.rows if _clean(row.full_name)]
    preview = await _preview(session, list(school_class.students), names)
    # Tên đăng nhập lấy **từ chính bảng xem trước**, kể cả khi một dòng trùng phía trên bị bỏ: tên
    # giáo viên vừa thấy trên màn 28b là tên được cấp (review pha 1-2). Một số đuôi bị nhảy không
    # sao; một tên khác với tên đã hiện thì có.
    keep = [
        row for row, sure in zip(preview.rows, confirmed, strict=True) if row.note == "new" or sure
    ]
    codes = await next_student_codes(session, len(keep))
    passwords = [new_password() for _ in keep]
    hashes = await asyncio.gather(*(hashed(one) for one in passwords))
    made: list[tuple[Student, str]] = []
    for row, code, password, digest in zip(keep, codes, passwords, hashes, strict=True):
        student = Student(
            class_id=class_id,
            full_name=row.full_name,
            student_code=code,
            username=row.username,
            password_hash=digest,
        )
        session.add(student)
        made.append((student, password))
    try:
        await session.commit()
    except IntegrityError as exc:
        await session.rollback()
        raise HTTPException(
            status_code=409, detail="Có giáo viên khác vừa thêm học sinh cùng lúc. Hãy bấm lại."
        ) from exc
    response.headers.update(_NO_STORE)
    return Created(
        accounts=[
            Account(
                student_id=student.id,
                full_name=student.full_name,
                student_code=student.student_code,
                username=student.username or "",
                password=password,
            )
            for student, password in made
        ],
        skipped=len(names) - len(keep),
    )


@router.post("/teacher/classes/{class_id}/roster/passwords.xlsx")
async def password_sheet(
    class_id: str,
    body: PasswordSheetIn,
    teacher: Teacher = Depends(current_teacher),
    session: AsyncSession = Depends(get_session),
) -> Response:
    """Bước 3: tệp mật khẩu `.xlsx`, dựng từ chính các dòng FE gửi lên.

    BE không có bản mật khẩu nào để so, nên không kiểm được các dòng là thật; điều nó kiểm được là
    lớp thuộc về người gọi. Tệp là thứ giáo viên đã thấy trên màn, không hơn.

    Raises:
        HTTPException: 404 khi lớp không phải của giáo viên này; 422 khi quá `MOST_ROWS` dòng.
    """
    school_class = await _owned_class(session, Asking.of(teacher), class_id)
    if not body.rows:
        raise HTTPException(status_code=422, detail="Không có dòng nào để xuất.")
    if len(body.rows) > MOST_ROWS:
        raise HTTPException(
            status_code=422, detail=f"Mỗi tệp mật khẩu có tối đa {MOST_ROWS} học sinh."
        )
    data = build_password_workbook(school_class.name, body.rows)
    return Response(
        content=data,
        media_type=_XLSX,
        headers=attachment(f"Mật khẩu - {school_class.name}.xlsx") | _NO_STORE,
    )


@router.post("/teacher/classes/{class_id}/students/{student_id}/password", response_model=Account)
async def reset_password(
    class_id: str,
    student_id: str,
    response: Response,
    teacher: Teacher = Depends(current_teacher),
    session: AsyncSession = Depends(get_session),
) -> Account:
    """*Đặt lại mật khẩu* (Figma 10b): mật khẩu mới, hiện một lần; mọi phiên cũ của em bị đóng.

    Một học sinh có từ trước việc 22 chưa có tên đăng nhập: lần đặt lại đầu tiên cấp luôn.

    Raises:
        HTTPException: 404 khi lớp không phải của giáo viên này, hay em không ở lớp ấy.

    Side effects:
        Ghi `password_hash` (và `username` nếu chưa có); xoá các dòng `sessions` của em; commit.
    """
    school_class = await _owned_class(session, Asking.of(teacher), class_id)
    student = next((one for one in school_class.students if one.id == student_id), None)
    if student is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy học sinh này trong lớp.")
    if student.username is None:
        student.username = (await usernames_for(session, [student.full_name]))[0]
    password = new_password()
    student.password_hash = await hashed(password)
    await session.execute(delete(AccountSession).where(AccountSession.student_id == student.id))
    try:
        await session.commit()
    except IntegrityError as exc:
        # Tên đăng nhập vừa cấp cho em có từ trước việc 22 bị một lần tạo khác lấy mất.
        await session.rollback()
        raise HTTPException(status_code=409, detail="Chưa đặt lại được, hãy bấm lại.") from exc
    response.headers.update(_NO_STORE)
    return Account(
        student_id=student.id,
        full_name=student.full_name,
        student_code=student.student_code,
        username=student.username,
        password=password,
    )
