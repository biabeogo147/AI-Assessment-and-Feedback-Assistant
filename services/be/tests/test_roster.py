"""Thêm học sinh vào lớp, tệp mật khẩu, đặt lại mật khẩu (việc 22, plan `sign-in-and-accounts`).

Ba học sinh demo có sẵn trong lớp 12A, trong đó *Nguyễn Minh Anh* (`anhnm`): đó là dòng trùng mà
các test dùng.
"""

from io import BytesIO
from urllib.parse import unquote

import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from be import db as db_module
from be.account_routes import router as account_router
from be.db import bind_sessions
from be.roster import router as roster_router
from be.seed import SAMPLE_STUDENTS, give_sample_sign_in, seed_if_empty
from be.student_routes import router as student_router
from be.teacher_reads import router as reads_router
from schema.ddl import prepare_schema
from schema.models import AccountSession, SchoolClass, Student, Teacher

TEACHER = {"X-Actor": "teacher:GV-001"}
STRANGER = {"X-Actor": "teacher:GV-002"}


@pytest_asyncio.fixture
async def stack():
    engine = create_async_engine("sqlite+aiosqlite://")
    await prepare_schema(engine)
    bind_sessions(engine)
    maker = async_sessionmaker(engine, expire_on_commit=False)
    async with maker() as session:
        await seed_if_empty(session)
        await give_sample_sign_in(session)
        session.add(Teacher(full_name="Thầy Nguyễn Văn B", teacher_code="GV-002"))
        await session.commit()
        class_id = await session.scalar(select(SchoolClass.id))
    app = FastAPI()
    for router in (account_router, roster_router, reads_router, student_router):
        app.include_router(router)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as http:
        yield http, maker, class_id
    await engine.dispose()
    db_module._SESSION_MAKER = None


def _xlsx(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.mark.asyncio
async def test_preview_marks_duplicates_and_writes_nothing(stack):
    http, maker, class_id = stack
    response = await http.post(
        f"/api/teacher/classes/{class_id}/roster/preview",
        headers=TEACHER,
        json={"names": ["  Phạm  Đức Anh ", "", "nguyễn minh anh", "Phạm Đức Anh"]},
    )
    assert response.status_code == 200
    assert response.json()["rows"] == [
        {"full_name": "Phạm Đức Anh", "username": "anhpd", "note": "new"},
        # Trùng em đã có trong lớp, không phân biệt hoa thường.
        {"full_name": "nguyễn minh anh", "username": "anhnm2", "note": "duplicate"},
        # Trùng dòng trên của chính loạt này.
        {"full_name": "Phạm Đức Anh", "username": "anhpd2", "note": "duplicate"},
    ]
    async with maker() as session:
        assert len((await session.scalars(select(Student))).all()) == 3


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("lop.csv", "﻿Họ và tên\nLê Kim Tính\n\nTrần Văn Cường,ghi chú\n".encode()),
        ("LOP.XLSX", _xlsx([["Họ tên"], ["Lê Kim Tính"], [None, "Trần Văn Cường"]])),
    ],
    ids=["csv", "xlsx"],
)
async def test_preview_from_file(stack, name, data):
    http, _, class_id = stack
    response = await http.post(
        f"/api/teacher/classes/{class_id}/roster/preview-file",
        headers=TEACHER,
        files={"file": (name, data)},
    )
    assert response.status_code == 200, response.text
    assert [(row["full_name"], row["username"]) for row in response.json()["rows"]] == [
        ("Lê Kim Tính", "tinhlk"),
        ("Trần Văn Cường", "cuongtv"),
    ]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("name", "data", "status", "detail"),
    [
        ("lop.txt", b"a", 422, "Chỉ nhận tệp .csv hoặc .xlsx."),
        ("lop.csv", "Lê Kim Tính".encode("utf-16"), 422, "Tệp .csv cần lưu ở mã UTF-8."),
        ("lop.xlsx", b"khong phai xlsx", 422, "Không đọc được tệp .xlsx này."),
        ("lop.csv", "Họ và tên\n".encode(), 422, "Danh sách chưa có họ tên nào."),
        ("lop.csv", b"a" * 1_000_001, 413, "Tệp danh sách lớn hơn 1 MB."),
    ],
    ids=["loai-tep", "khong-utf8", "xlsx-hong", "chi-tieu-de", "qua-lon"],
)
async def test_preview_file_refusals(stack, name, data, status, detail):
    http, _, class_id = stack
    response = await http.post(
        f"/api/teacher/classes/{class_id}/roster/preview-file",
        headers=TEACHER,
        files={"file": (name, data)},
    )
    assert response.status_code == status
    assert response.json()["detail"] == detail


@pytest.mark.asyncio
async def test_create_skips_unconfirmed_duplicates(stack):
    http, maker, class_id = stack
    response = await http.post(
        f"/api/teacher/classes/{class_id}/roster",
        headers=TEACHER,
        json={
            "rows": [
                {"full_name": "Lê Kim Tính"},
                {"full_name": "Nguyễn Minh Anh", "confirmed": False},
                {"full_name": "Lê Thu Hà", "confirmed": True},
            ]
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["skipped"] == 1
    made = body["accounts"]
    assert [(one["full_name"], one["username"], one["student_code"]) for one in made] == [
        ("Lê Kim Tính", "tinhlk", "1"),
        ("Lê Thu Hà", "halt2", "2"),
    ]
    async with maker() as session:
        stored = await session.scalar(select(Student).where(Student.username == "tinhlk"))
        assert stored.class_id == class_id
        assert made[0]["password"] not in stored.password_hash

    # Em vừa tạo đăng nhập được bằng chính mật khẩu vừa hiện.
    signed = await http.post(
        "/api/sign-in", json={"username": "tinhlk", "password": made[0]["password"]}
    )
    assert signed.status_code == 204
    assert (await http.get("/api/me")).json()["student_code"] == "1"


@pytest.mark.asyncio
async def test_roster_is_the_owners_only(stack):
    http, _, class_id = stack
    for path, kwargs in (
        ("roster/preview", {"json": {"names": ["A"]}}),
        ("roster", {"json": {"rows": [{"full_name": "A"}]}}),
        (
            "roster/passwords.xlsx",
            {"json": {"rows": [{"full_name": "A", "username": "a", "password": "p"}]}},
        ),
    ):
        response = await http.post(
            f"/api/teacher/classes/{class_id}/{path}", headers=STRANGER, **kwargs
        )
        assert response.status_code == 404, path


@pytest.mark.asyncio
async def test_password_sheet_is_what_was_sent(stack):
    http, _, class_id = stack
    rows = [
        {"full_name": "Lê Kim Tính", "username": "tinhlk", "password": "k7m-pq4-xd"},
        {"full_name": '=HYPERLINK("x")', "username": "x", "password": "abc-def-gh"},
    ]
    response = await http.post(
        f"/api/teacher/classes/{class_id}/roster/passwords.xlsx",
        headers=TEACHER,
        json={"rows": rows},
    )
    assert response.status_code == 200
    assert unquote(response.headers["content-disposition"].split("UTF-8''")[1]) == (
        "Mật khẩu - 12A.xlsx"
    )
    sheet = load_workbook(BytesIO(response.content)).active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values == [
        ["Họ và tên", "Tên đăng nhập", "Mật khẩu"],
        ["Lê Kim Tính", "tinhlk", "k7m-pq4-xd"],
        ['=HYPERLINK("x")', "x", "abc-def-gh"],
    ]
    assert sheet["A3"].data_type == "s"


@pytest.mark.asyncio
async def test_reset_password_shows_once_and_ends_old_sessions(stack):
    http, maker, class_id = stack
    code, username, old = SAMPLE_STUDENTS[0]
    await http.post("/api/sign-in", json={"username": username, "password": old})
    async with maker() as session:
        student = await session.scalar(select(Student).where(Student.student_code == code))
    response = await http.post(
        f"/api/teacher/classes/{class_id}/students/{student.id}/password", headers=TEACHER
    )
    assert response.status_code == 200
    fresh = response.json()["password"]
    assert response.json()["username"] == username
    # Phiên cũ của em đã đóng ...
    assert (await http.get("/api/me")).status_code == 401
    async with maker() as session:
        assert await session.scalar(select(AccountSession.id)) is None
    # ... mật khẩu cũ hết tác dụng, mật khẩu mới thì vào được.
    assert (
        await http.post("/api/sign-in", json={"username": username, "password": old})
    ).status_code == 401
    assert (
        await http.post("/api/sign-in", json={"username": username, "password": fresh})
    ).status_code == 204


@pytest.mark.asyncio
async def test_reset_gives_a_username_to_older_students(stack):
    http, maker, class_id = stack
    async with maker() as session:
        older = Student(class_id=class_id, full_name="Đỗ Văn Hùng", student_code="HS2026-1300")
        session.add(older)
        await session.commit()
    detail = (await http.get(f"/api/teacher/classes/{class_id}", headers=TEACHER)).json()
    assert {one["student_code"]: one["username"] for one in detail["students"]}[
        "HS2026-1300"
    ] is None

    response = await http.post(
        f"/api/teacher/classes/{class_id}/students/{older.id}/password", headers=TEACHER
    )
    assert response.json()["username"] == "hungdv"
    detail = (await http.get(f"/api/teacher/classes/{class_id}", headers=TEACHER)).json()
    assert {one["student_code"]: one["username"] for one in detail["students"]}[
        "HS2026-1300"
    ] == "hungdv"


@pytest.mark.asyncio
async def test_reset_refuses_a_student_of_another_class(stack):
    http, maker, class_id = stack
    async with maker() as session:
        other = SchoolClass(
            teacher_id=await session.scalar(
                select(Teacher.id).where(Teacher.teacher_code == "GV-001")
            ),
            name="12B",
        )
        session.add(other)
        await session.flush()
        stray = Student(class_id=other.id, full_name="Em Lạ", student_code="HS2026-9999")
        session.add(stray)
        await session.commit()
    response = await http.post(
        f"/api/teacher/classes/{class_id}/students/{stray.id}/password", headers=TEACHER
    )
    assert response.status_code == 404
    assert response.json()["detail"] == "Không tìm thấy học sinh này trong lớp."


@pytest.mark.asyncio
async def test_codes_sort_by_value_in_class_detail(stack):
    http, _, class_id = stack
    names = [f"Học Sinh {n}" for n in range(11)]
    await http.post(
        f"/api/teacher/classes/{class_id}/roster",
        headers=TEACHER,
        json={"rows": [{"full_name": name} for name in names]},
    )
    detail = (await http.get(f"/api/teacher/classes/{class_id}", headers=TEACHER)).json()
    codes = [one["student_code"] for one in detail["students"]]
    assert codes == [*(str(n) for n in range(1, 12)), "HS2026-1204", "HS2026-1205", "HS2026-1206"]


@pytest.mark.asyncio
async def test_created_usernames_are_the_previewed_ones(stack):
    """Dòng trùng bị bỏ ở trên không làm tên của dòng dưới lệch khỏi bảng xem trước (review)."""
    http, _, class_id = stack
    names = ["Nguyễn Minh Anh", "Nguyễn Mạnh Anh"]
    preview = await http.post(
        f"/api/teacher/classes/{class_id}/roster/preview", headers=TEACHER, json={"names": names}
    )
    shown = [row["username"] for row in preview.json()["rows"]]
    assert shown == ["anhnm2", "anhnm3"]
    created = await http.post(
        f"/api/teacher/classes/{class_id}/roster",
        headers=TEACHER,
        json={"rows": [{"full_name": name} for name in names]},
    )
    assert [one["username"] for one in created.json()["accounts"]] == ["anhnm3"]
    assert created.headers["cache-control"] == "no-store"


@pytest.mark.asyncio
async def test_a_very_long_name_still_fits_the_username_column(stack):
    http, _, class_id = stack
    created = await http.post(
        f"/api/teacher/classes/{class_id}/roster",
        headers=TEACHER,
        json={"rows": [{"full_name": "A" * 120}, {"full_name": "A" * 120, "confirmed": True}]},
    )
    usernames = [one["username"] for one in created.json()["accounts"]]
    assert usernames == ["a" * 56, "a" * 56 + "2"]


@pytest.mark.asyncio
async def test_a_csv_cell_past_the_field_limit_is_a_refusal(stack):
    http, _, class_id = stack
    response = await http.post(
        f"/api/teacher/classes/{class_id}/roster/preview-file",
        headers=TEACHER,
        files={"file": ("lop.csv", b'"' + b"a" * 200_000 + b'"\n')},
    )
    assert response.status_code == 422
    assert response.json()["detail"] == "Không đọc được tệp .csv này."


@pytest.mark.asyncio
async def test_password_responses_are_not_cached(stack):
    http, maker, class_id = stack
    async with maker() as session:
        student_id = await session.scalar(select(Student.id).where(Student.username == "anhnm"))
    reset = await http.post(
        f"/api/teacher/classes/{class_id}/students/{student_id}/password", headers=TEACHER
    )
    sheet = await http.post(
        f"/api/teacher/classes/{class_id}/roster/passwords.xlsx",
        headers=TEACHER,
        json={"rows": [{"full_name": "A", "username": "a", "password": "p"}]},
    )
    assert reset.headers["cache-control"] == sheet.headers["cache-control"] == "no-store"
